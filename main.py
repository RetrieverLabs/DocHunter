import os
import re
import csv
import zipfile
import io
import multiprocessing as mp
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Dict

# 3rd-party libraries
import pypdf
from pptx import Presentation
import openpyxl
import xlrd
from bs4 import BeautifulSoup
from PIL import Image
import pytesseract

SUPPORTED_TEXT_EXTENSIONS = {
    ".txt", ".csv", ".json", ".xml", ".docx",
    ".doc", ".dot", ".html", ".md", ".pdf",
    ".png", ".pptx", ".rtf", ".tif", ".xls",
    ".xlsx", ".zip"
}

DEFAULT_WORKERS = max(1, mp.cpu_count() // 2)

PII_PATTERNS: Dict[str, str] = {
    "Credit Card": r"\b(?:4(?:[\s-]*\d){12}(?:(?:[\s-]*\d){3})?|(?:5[1-5]|2[2-7][2\-0])(?:[\s-]*\d){14}|3[47](?:[\s-]*\d){13}|6(?:011|5|4[4-9]|22)(?:[\s-]*\d){12,15})\b",
    "US SSN": r"\b\d{3}[\s-]+\d{2}[\s-]+\d{4}\b",
    "Canadian SIN": r"\b(?:\d{9}|\d{3}-\d{3}-\d{3}|\d{3}\s\d{3}\s\d{3})\b",
    "US/CA Drivers License": r"\b([A-Z](?:[\s-]*\d){14}|[A-Z](?:[\s-]*\d){12}|\d{7,9}|[A-Z]\d{7}\b)"
}

class Progress:
    def __init__(self, total: int):
        self.total = total
        self.done = 0
        self.matches = 0

    def update(self, match_count: int):
        self.done += 1
        self.matches += match_count
        self.render()

    def render(self):
        if self.total == 0:
            return
        pct = (self.done / self.total) * 100
        bar_len = 30
        filled = int(bar_len * self.done / self.total)
        bar = "█" * filled + "░" * (bar_len - filled)
        print(
            f"\r[{bar}] {pct:5.1f}% "
            f"{self.done}/{self.total} files "
            f"| matches: {self.matches}",
            end=""
        )

class MatchResult:
    def __init__(self, file_path: str, extension: str, pii_type: str, line_number: int, context: str):
        self.file_path = file_path
        self.extension = extension
        self.pii_type = pii_type
        self.line_number = line_number
        self.context = context

def extract_strings_from_binary(content: bytes) -> List[str]:
    """Fallback text extraction for binary legacy files (.doc, .dot, .rtf)."""
    text = re.sub(r'[^\x20-\x7E\n\r\t]', ' ', content.decode('latin-1', errors='ignore'))
    return [line.strip() for line in text.splitlines() if len(line.strip()) > 4]

def extract_from_stream(stream: io.BytesIO, ext: str) -> List[str]:
    """Extract text lines safely from a byte stream based on extension."""
    lines = []
    ext = ext.lower()

    try:
        if ext in {".txt", ".csv", ".json", ".xml", ".md"}:
            lines = stream.read().decode("utf-8", errors="ignore").splitlines()

        elif ext == ".html":
            soup = BeautifulSoup(stream.read(), "html.parser")
            lines = soup.get_text(separator="\n").splitlines()

        elif ext == ".docx":
            with zipfile.ZipFile(stream) as docx:
                xml_content = docx.read('word/document.xml')
                root = ET.fromstring(xml_content)
                namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for paragraph in root.findall('.//w:p', namespaces):
                    text_pieces = [node.text for node in paragraph.findall('.//w:t', namespaces) if node.text]
                    if text_pieces:
                        lines.append("".join(text_pieces))

        elif ext == ".pdf":
            reader = pypdf.PdfReader(stream)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    lines.extend(text.splitlines())

        elif ext == ".pptx":
            prs = Presentation(stream)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        lines.extend([p.text for p in shape.text_frame.paragraphs if p.text])
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_text:
                                lines.append(" | ".join(row_text))

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(stream, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_str = " ".join([str(cell) for cell in row if cell is not None])
                    if row_str:
                        lines.append(row_str)

        elif ext == ".xls":
            wb = xlrd.open_workbook(file_contents=stream.read())
            for sheet in wb.sheets():
                for row_idx in range(sheet.nrows):
                    row_vals = sheet.row_values(row_idx)
                    row_str = " ".join([str(val) for val in row_vals if val])
                    if row_str:
                        lines.append(row_str)

        elif ext in {".png", ".tif"}:
            img = Image.open(stream)
            text = pytesseract.image_to_string(img)
            lines = text.splitlines()

        elif ext in {".doc", ".dot", ".rtf"}:
            lines = extract_strings_from_binary(stream.read())

    except Exception:
        # Prevents unreadable or corrupt individual files from halting execution
        pass

    return lines

def process_archive(zip_stream: io.BytesIO, archive_path: str) -> List[MatchResult]:
    """Recursively process supported files inside a ZIP archive."""
    results = []
    try:
        with zipfile.ZipFile(zip_stream) as zf:
            for member in zf.infolist():
                if member.is_dir():
                    continue
                member_ext = Path(member.filename).suffix.lower()
                if member_ext in SUPPORTED_TEXT_EXTENSIONS:
                    with zf.open(member) as f:
                        file_bytes = io.BytesIO(f.read())
                        virtual_path = f"{archive_path}::{member.filename}"
                        if member_ext == ".zip":
                            results.extend(process_archive(file_bytes, virtual_path))
                        else:
                            lines = extract_from_stream(file_bytes, member_ext)
                            results.extend(scan_lines(lines, virtual_path, member_ext))
    except Exception:
        pass
    return results

def scan_lines(lines: List[str], file_path: str, extension: str) -> List[MatchResult]:
    """Scan extracted text lines against regex PII patterns."""
    results = []
    for i, line in enumerate(lines, start=1):
        clean_line = line.strip()
        if not clean_line:
            continue

        for pii_label, pattern in PII_PATTERNS.items():
            if re.search(pattern, clean_line):
                results.append(
                    MatchResult(
                        file_path=file_path,
                        extension=extension,
                        pii_type=pii_label,
                        line_number=i,
                        context=clean_line[:300]
                    )
                )
                break
    return results

def search_worker(file_path: str) -> List[MatchResult]:
    path = Path(file_path)
    ext = path.suffix.lower()

    try:
        with open(path, "rb") as f:
            stream = io.BytesIO(f.read())

        if ext == ".zip":
            return process_archive(stream, str(path))
        else:
            lines = extract_from_stream(stream, ext)
            return scan_lines(lines, str(path), ext)
    except Exception:
        return []

def find_files(root: Path) -> List[Path]:
    files = []
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_TEXT_EXTENSIONS:
            files.append(path)
    return files

def main():
    target_directory = Path(__file__).resolve().parent
    
    print("==========================================")
    print("🛡️  DocHunter: Clean PII Scanner        🛡️")
    print("==========================================")
    
    files = find_files(target_directory)
    print(f"🔍 Analyzing {len(files)} target files for PII leaks using {DEFAULT_WORKERS} workers...\n")
    
    results: List[MatchResult] = []
    file_strings = [str(f.resolve()) for f in files]

    progress = Progress(total=len(file_strings))
    progress.render()

    with mp.Pool(processes=DEFAULT_WORKERS) as pool:
        for chunk_results in pool.imap_unordered(search_worker, file_strings):
            results.extend(chunk_results)
            progress.update(match_count=len(chunk_results))

    print()

    output_path = "results.csv"
    headers = ["File Path", "Extension", "Detected PII Type", "Line Number", "Text Excerpt Preview"]
    try:
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for r in results:
                writer.writerow([r.file_path, r.extension, r.pii_type, r.line_number, r.context])
        print(f"\n🎉 Audit complete! Report saved safely to: {output_path}")
    except Exception as e:
        print(f"Error compiling output file: {e}")

if __name__ == "__main__":
    if os.name != 'nt':
        try:
            mp.set_start_method("fork", force=True)
        except RuntimeError:
            pass
    else:
        try:
            mp.set_start_method("spawn", force=True)
        except RuntimeError:
            pass
        
    try:
        main()
    finally:
        # Keeps the console window open when double-clicking the file directly
        input("\nPress Enter to exit...")
